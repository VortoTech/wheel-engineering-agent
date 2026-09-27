"""Turn the raw factory order folders into a privacy-scrubbed case set.

    uv run --with openpyxl --with xlrd --with olefile python experiments/real-orders/ingest.py \
        --raw ~/personal_project/asset [--out runs/real-orders]

Raw folder layout (one folder per order, name holding the design code HNxxx / Mxx): three renders
(front, oblique, back; two wheels side by side when the order has two sizes, a spec text block in a
corner), the factory order sheet (*确认单*.xlsx), the customer request (*.xls), Parasolid x_t per
size, DWG drawings.

Output, per case (case-01, ...) and size (d20w10.5, ...):
    OUT/case-01/order.json                      engineering fields of the order sheets
    OUT/case-01/d20w10.5/{front,oblique,back}.jpg   one wheel per image, text block removed, no
                                                    metadata (cap emblems blurred with --blur-logos)
    OUT/case-01/d20w10.5/spec.json              the six key specs, source "order_sheet"
    OUT/case-01/d20w10.5/model.x_t              the CAD with its private header fields redacted
    OUT/case-01/d20w10.5/truth.json             dimensions read from the CAD (parasolid_xt.wheel_truth)
DWG drawings are not copied (binary headers are not scrubbed). The case map (real folder -> case)
and the deny list of private terms found in the sheets and headers go to ~/.wheelcam/, never to OUT
or the repo; OUT is scanned against the deny list at the end.
"""
import argparse
import json
import re
import sys
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "services"))

from wheelcam.parasolid_xt import read_xt, wheel_truth  # noqa: E402
from wheelcam.privacy import clean_image, order_from_rows, private_terms, scan, scrub_xt  # noqa: E402

CODE_RE = re.compile(r"\b(HN\d+|M\d+)\b", re.I)
XT_SIZE_RE = re.compile(r"(\d{2})[Xx]?(\d{1,2}(?:\.\d+)?)-(\d+)[-Xx](\d+(?:\.\d+)?)-\(?(-?\d+)\)?-(\d+(?:\.\d+)?)\.x_t$", re.I)
PRIVATE_DIR = Path("~/.wheelcam").expanduser()
MAX_SIDE = 2000


def sheet_rows(path: Path):
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        return [list(r) for r in openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True)]
    import xlrd
    ws = xlrd.open_workbook(path).sheet_by_index(0)
    return [ws.row_values(r) for r in range(ws.nrows)]


def sheet_private_terms(path: Path) -> set:
    """Names and codes in a sheet and its document properties (for the deny list)."""
    terms = set()
    rows = sheet_rows(path)
    for row in rows:
        cells = [str(c).strip() for c in row if c not in (None, "")]
        for i, c in enumerate(cells[:-1]):
            if c.replace(" ", "").rstrip("：:") in ("客户名称", "审核", "制表", "联系电话", "收货地址", "客户代码"):
                value = cells[i + 1]
                terms.add(value)
                if re.fullmatch(r"[A-Z]{2,3}", value):
                    terms.add(value + "-")                          # a customer code, as engraved
        for c in cells:
            terms.update(m.group(0) for m in re.finditer(r"\b[A-Z]{2,3}-(?:HN|M)\d+", c))
            if "锻造轮毂客户确认单" in c and c.index("锻造") > 0:
                terms.add(c[:c.index("锻造")])                      # the factory name in the title
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        props = openpyxl.load_workbook(path).properties
        terms.update(v for v in (props.creator, props.lastModifiedBy) if v)
    else:
        import olefile
        with olefile.OleFileIO(str(path)) as ole:
            meta = ole.get_metadata()
            terms.update(v.decode("latin-1") for v in (meta.author, meta.last_saved_by) if v)
    return {t.strip() for t in terms if keep_term(t)}


def xt_private_terms(text: str) -> set:
    """Machine and account names in a Parasolid header (the folders of FILE are not personal)."""
    head = text.split("**END_OF_HEADER")[0]
    terms = set()
    for key in ("MC_ID", "USER"):
        m = re.search(rf"(?m)^{key}=([^;]*);", head)
        if m and m.group(1).strip() not in ("", "^_", "unknown", "admin"):
            terms.add(m.group(1).strip())
    return terms


def keep_term(t: str) -> bool:
    """A deny-list term specific enough to scan for: ASCII at least 4 characters, CJK at least 2,
    not a sheet label."""
    t = t.strip()
    if not t or t.endswith(("：", ":")) or "：" in t:
        return False
    return len(t) >= (2 if re.search(r"[\u4e00-\u9fff]", t) else 4)


def wheels(image: np.ndarray):
    """[(x0, x1, mask)] per wheel side by side: the largest foreground part of each half (or of the
    whole image when it is square); the spec text block and other specks are left out."""
    from scipy.ndimage import binary_fill_holes, label
    h, w = image.shape[:2]
    spans = [(0, w // 2), (w // 2, w)] if w > 1.5 * h else [(0, w)]
    out = []
    for x0, x1 in spans:
        fg = image[:, x0:x1, :3].min(axis=2) < .92
        lab, n = label(fg)
        if not n:
            continue
        sizes = np.bincount(lab.ravel())[1:]
        big = sizes.argmax() + 1
        # the wheel may fall apart into a few large parts (thin spokes on a light finish)
        keep = [i + 1 for i, s in enumerate(sizes) if s > .02 * sizes.max()]
        mask = np.isin(lab, keep)
        ys, xs = np.nonzero(lab == big)
        cx, cy, r = xs.mean(), ys.mean(), .6 * max(np.ptp(xs), np.ptp(ys))
        yy, xx = np.mgrid[:h, :x1 - x0]
        mask &= (xx - cx) ** 2 + (yy - cy) ** 2 < r * r            # drop the text block in the corner
        out.append((x0, x1, binary_fill_holes(mask)))
    return out


def logo_circles(a: np.ndarray, min_px=30):
    """Circles over coloured emblems (a car maker's roundel on the cap): clusters of saturated blue
    or red pixels. The finishes (black, bronze, silver, grey) are not blue or red."""
    from scipy.ndimage import binary_dilation, label
    import colorsys
    rgb = a[..., :3]
    mx, mn = rgb.max(axis=2), rgb.min(axis=2)
    sat = (mx - mn) / np.maximum(mx, 1e-6)
    r, g, b = rgb[..., 0], rgb[..., 1], rgb[..., 2]
    blue = (b > r + .12) & (b > g + .05) & (sat > .35) & (mx > .2)
    red = (r > g + .25) & (r > b + .25) & (sat > .5) & (mx > .3)
    lab, n = label(binary_dilation(blue | red, iterations=6))
    out = []
    for k in range(1, n + 1):
        ys, xs = np.nonzero(lab == k)
        if len(xs) < min_px:
            continue
        rad = .5 * max(np.ptp(xs), np.ptp(ys))
        out.append((float(xs.mean()), float(ys.mean()), 1.8 * rad + 6))
    return out


def split_render(src: Path, view: str, dst_by_index, cap_frac=.075, blur_logos=False):
    """Write each wheel of `src` to dst_by_index(i) (i = left to right); returns the wheel pixel
    diameters. With blur_logos the centre cap and coloured emblems are blurred (for showing)."""
    from PIL import Image
    im = Image.open(src).convert("RGB")
    scale = MAX_SIDE / max(im.size) * (2 if im.size[0] > 1.5 * im.size[1] else 1)
    if scale < 1:
        im = im.resize((round(im.size[0] * scale), round(im.size[1] * scale)), Image.LANCZOS)
    a = np.asarray(im, float) / 255
    sizes = []
    for i, (x0, x1, mask) in enumerate(wheels(a)):
        part = a[:, x0:x1].copy()
        part[~mask] = 1.0                                           # white outside the wheel
        ys, xs = np.nonzero(mask)
        pad = int(.04 * max(np.ptp(xs), np.ptp(ys)))
        y0, y1 = max(0, ys.min() - pad), min(part.shape[0], ys.max() + pad)
        xa, xb = max(0, xs.min() - pad), min(part.shape[1], xs.max() + pad)
        part = part[y0:y1, xa:xb]
        d = float(max(np.ptp(xs), np.ptp(ys)))
        sizes.append(np.ptp(ys) if view == "front" else d)
        blur = logo_circles(part) if blur_logos else []
        if blur_logos and view == "front":                          # the whole centre cap, straight on
            try:
                from wheelcam.wheel_skill import _front_rim_hub
                found = _front_rim_hub(part)
                if found:
                    rim, hub = found
                    radius = float(np.max(np.linalg.norm(np.asarray(rim) - np.asarray(hub), axis=1)))
                    blur.append((hub[0], hub[1], cap_frac * 2 * radius))
            except Exception:
                pass
        tmp = Image.fromarray((part * 255).astype("uint8"))
        dst = dst_by_index(i)
        dst.parent.mkdir(parents=True, exist_ok=True)
        tmp.save(dst.with_suffix(".tmp.png"))
        clean_image(dst.with_suffix(".tmp.png"), dst, blur=blur)
        dst.with_suffix(".tmp.png").unlink()
    return sizes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--out", default="runs/real-orders")
    ap.add_argument("--blur-logos", action="store_true", help="blur cap emblems (for a public showing)")
    a = ap.parse_args()
    raw, out = Path(a.raw).expanduser(), Path(a.out)
    folders = {}
    for d in sorted(p for p in raw.rglob("*") if p.is_dir()):
        m = CODE_RE.search(d.name)
        files = [f for f in d.iterdir() if f.is_file()]
        if m and any(f.suffix.lower() in (".jpg", ".png", ".x_t") for f in files):
            folders[m.group(1).upper()] = d
    PRIVATE_DIR.mkdir(exist_ok=True)
    map_file = PRIVATE_DIR / "case_map.json"
    case_map = json.loads(map_file.read_text()) if map_file.exists() else {}
    for code in sorted(folders):
        case_map.setdefault(code, f"case-{len(case_map) + 1:02d}")
    map_file.write_text(json.dumps(case_map, ensure_ascii=False, indent=1))
    terms = set(private_terms())
    summary = []
    for code, folder in sorted(folders.items(), key=lambda kv: case_map[kv[0]]):
        case = case_map[code]
        cdir = out / case
        files = sorted(f for f in folder.iterdir() if f.is_file())
        sheets = [f for f in files if f.suffix.lower() in (".xlsx", ".xls")]
        factory = [f for f in sheets if f.suffix.lower() == ".xlsx"]
        order = order_from_rows(sheet_rows(factory[0])) if factory else {"variants": []}
        if not order["variants"]:
            order = order_from_rows(sheet_rows(sheets[0])) if sheets else order
        for f in sheets:
            terms |= sheet_private_terms(f)
        variants = sorted(order["variants"], key=lambda v: (v["diameter_in"], v["width_in"]))
        tag = lambda v: f"d{v['diameter_in']}w{v['width_in']:g}"
        cdir.mkdir(parents=True, exist_ok=True)
        (cdir / "order.json").write_text(json.dumps({**order, "variants": variants, "source": "order_sheet"},
                                                    ensure_ascii=False, indent=1))
        renders = sorted((f for f in files if f.suffix.lower() in (".jpg", ".png")),
                         key=lambda f: int(re.findall(r"(\d+)", f.stem)[0]))
        order_by_size = None
        for view, src in zip(("front", "oblique", "back"), renders):
            dsts = {}
            sizes = split_render(src, view, lambda i: dsts.setdefault(i, cdir / f"_w{i}" / f"{view}.jpg"),
                                 blur_logos=a.blur_logos)
            if view == "front":
                order_by_size = list(np.argsort(sizes))                # smaller wheel = smaller size
        for i, v in enumerate(variants):
            vdir = cdir / tag(v)
            src_dir = cdir / f"_w{order_by_size[i] if order_by_size and i < len(order_by_size) else i}"
            if src_dir.exists():
                vdir.mkdir(exist_ok=True)
                for f in src_dir.iterdir():
                    f.replace(vdir / f.name)
            spec = {k: v[k] for k in ("diameter_in", "width_in", "pcd_mm", "bolts", "center_bore_mm", "et_mm") if k in v}
            vdir.mkdir(exist_ok=True)
            (vdir / "spec.json").write_text(json.dumps({"spec": spec, "source": "order_sheet",
                                                        "hole_form": v.get("hole_form")}, indent=1))
        for d in cdir.glob("_w*"):
            for f in d.iterdir():
                f.unlink()
            d.rmdir()
        for f in (f for f in files if f.suffix.lower() == ".x_t"):
            text = f.read_text(encoding="latin-1")
            terms |= {t for t in xt_private_terms(text) if keep_term(t)}
            m = XT_SIZE_RE.search(f.name)
            match = [v for v in variants if m and v["diameter_in"] == int(m.group(1))
                     and abs(v["width_in"] - float(m.group(2))) < 1e-6] if m else []
            if match:
                (cdir / tag(match[0]) / "model.x_t").write_text(scrub_xt(text), encoding="latin-1")
                truth = wheel_truth(read_xt(text))
                (cdir / tag(match[0]) / "truth.json").write_text(json.dumps({**truth, "source": "factory_cad_x_t"}, indent=1))
        summary.append({"case": case, "sizes": [tag(v) for v in variants],
                        "renders": len(renders), "cad": sorted(p.parent.name for p in cdir.glob("*/model.x_t"))})
    terms_file = PRIVATE_DIR / "private_terms.txt"
    old = set(private_terms(terms_file))
    terms_file.write_text("\n".join(sorted(old | terms)) + "\n", encoding="utf-8")
    (out / "cases.json").write_text(json.dumps(summary, indent=1))
    leaks = scan([out], sorted(old | terms))
    for s in summary:
        print(json.dumps(s))
    print(f"deny list: {len(old | terms)} terms ({terms_file}); leaks in output: {len(leaks)}")
    if leaks:
        print(json.dumps({Path(k).relative_to(out).as_posix(): v for k, v in leaks.items()}, ensure_ascii=False, indent=1))
        sys.exit(1)


if __name__ == "__main__":
    main()
